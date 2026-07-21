#!/usr/bin/env python3
"""Reconcile scraped_people.json against people-roster.xlsx.

Matches each scraped profile to a roster row (fuzzy on normalized name and
surname+first-initial), reports matches/misses. With --write it merges the
scraped fields into the roster (adding Title/Bio/Email/Scholar/ORCID/Twitter/
Photo columns and new rows for people absent from the roster) and emits a photo
download manifest + a wix-slug->profile-slug map for the 404 router.
"""
import json, re, sys, os
import openpyxl

HERE = os.path.dirname(__file__)
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
ROSTER = os.path.join(ROOT, "data", "people-roster.xlsx")
SCRAPE = os.path.join(HERE, "scraped_people.json")

# wix-slug -> exact roster name, for cases fuzzy matching can't resolve
OVERRIDES = {
    "edwardbuckler": "Edward S. Buckler",
    "cintaromay": "M. Cinta Romay",
    "saralarsson": "Sara Larsson",
    "mikegore": "Michael Gore",
    "joegage": "Joseph Gage",
    "mejia-guerra": "Maria Katherine Mejia-Guerra",
    "rodgers-melnick": "Eli Rodgers-Melnick",
    "germanocosta-ento": "Germano Costa-Neto",
    "irievroh": "Irie Vroh Bi",
    "patrickjbrown": "Patrick Brown",
    "dallaskroon": "Dallas E. Kroon",
    "niranjani-gnanapragasm": "Niranjani Gnanapragasam",
    "xiaoleiliu": "Xiaolei Liu",
    "carloslignecalderonvazquez": "Carlos Ligne Calderon Vazquez",
    "mei-hsiusu": "Mei-Hsiu Su",
    "ramsharma": "Ram Sharma",
    "kentashirasawa": "Kenta Shirasawa",
}

def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())

def loose(s):
    toks = re.sub(r"[^a-z ]", " ", (s or "").lower()).split()
    toks = [t for t in toks if len(t) > 1]  # drop initials
    if not toks: return ""
    return toks[-1] + toks[0][0]  # surname + first initial

def slugify(s):
    return re.sub(r"^-|-$", "", re.sub(r"[^a-z0-9]+", "-", (s or "").lower()))

def main():
    write = "--write" in sys.argv
    wb = openpyxl.load_workbook(ROSTER)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    name_col = headers.index("Name")
    rows = []
    for r in ws.iter_rows(min_row=2):
        nm = r[name_col].value
        if nm: rows.append((r, str(nm).strip()))
    by_norm, by_loose = {}, {}
    for r, nm in rows:
        by_norm.setdefault(norm(nm), []).append((r, nm))
        by_loose.setdefault(loose(nm), []).append((r, nm))

    scraped = json.load(open(SCRAPE))
    matched, unmatched = {}, {}
    for wixslug, rec in scraped.items():
        nm = rec.get("name", "")
        target = None
        if wixslug in OVERRIDES:
            key = norm(OVERRIDES[wixslug])
            if key in by_norm: target = by_norm[key][0][1]
        if not target and norm(nm) in by_norm and len(by_norm[norm(nm)]) == 1:
            target = by_norm[norm(nm)][0][1]
        if not target and loose(nm) in by_loose and len(by_loose[loose(nm)]) == 1:
            target = by_loose[loose(nm)][0][1]
        if target:
            matched[wixslug] = target
        else:
            unmatched[wixslug] = nm

    print(f"scraped={len(scraped)}  matched={len(matched)}  new(unmatched)={len(unmatched)}")
    print("\n--- NEW (will be added as roster rows) ---")
    for w, nm in sorted(unmatched.items(), key=lambda x: x[1]):
        print(f"  {nm}   [{w}]")

    if not write:
        print("\n(dry run — pass --write to apply)")
        return

    # ---- ensure new columns exist ----
    NEWCOLS = ["Title", "Bio", "Email", "Scholar", "ORCID", "Twitter", "Photo"]
    for c in NEWCOLS:
        if c not in headers:
            headers.append(c)
            ws.cell(row=1, column=len(headers), value=c)
    col = {h: i + 1 for i, h in enumerate(headers)}  # 1-based

    def role_from_title(t):
        t = (t or "").lower()
        if "postdoc" in t or "post doc" in t or "post-doc" in t: return "postdoc"
        if "grad" in t or "phd" in t or "ph.d" in t or "doctoral" in t: return "grad"
        if "visit" in t: return "visiting"
        if "undergrad" in t or "intern" in t or "student" in t: return "undergrad"
        return "staff"

    def setcell(rownum, header, val, only_if_empty=False):
        if val is None: return
        cell = ws.cell(row=rownum, column=col[header])
        if only_if_empty and cell.value not in (None, ""): return
        cell.value = val

    # index roster row numbers by name
    rowByName = {nm: r[0].row for r, nm in rows}
    photo_manifest = []   # (slug, url)
    router_map = {}       # wixslug -> profile slug

    def apply(rownum, canonical_name, rec):
        slug = slugify(canonical_name)
        setcell(rownum, "Title", rec.get("title"))
        setcell(rownum, "Bio", rec.get("bio"))
        setcell(rownum, "Email", rec.get("email"))
        setcell(rownum, "Scholar", rec.get("scholar"))
        setcell(rownum, "ORCID", rec.get("orcid"))
        setcell(rownum, "Twitter", rec.get("twitter"))
        setcell(rownum, "Website", rec.get("website"), only_if_empty=True)
        setcell(rownum, "LinkedIn", rec.get("linkedin"), only_if_empty=True)
        if rec.get("photo"):
            setcell(rownum, "Photo", slug + ".jpg")
            photo_manifest.append((slug, rec["photo"]))
        return slug

    for wixslug, rec in scraped.items():
        if wixslug in matched:
            canonical = matched[wixslug]
            apply(rowByName[canonical], canonical, rec)
        else:
            canonical = rec["name"]
            ws.append({col["Name"]: canonical, col["Role"]: role_from_title(rec.get("title"))})
            apply(ws.max_row, canonical, rec)
        router_map[wixslug] = slugify(canonical)

    wb.save(ROSTER)
    with open(os.path.join(HERE, "photo_manifest.tsv"), "w") as f:
        for slug, url in photo_manifest:
            f.write(f"{slug}\t{url}\n")
    json.dump(router_map, open(os.path.join(HERE, "router_person_map.json"), "w"),
              indent=1, ensure_ascii=False)
    print(f"\nwrote roster (+{len(unmatched)} rows), "
          f"{len(photo_manifest)} photos to fetch, {len(router_map)} router entries")

if __name__ == "__main__":
    main()

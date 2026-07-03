#!/usr/bin/env python3
# Seed data/people-roster.xlsx from data/people.tsv (the build_people.py output).
#
# Run this ONLY to (re)create the roster spreadsheet from the generated seed —
# it OVERWRITES the xlsx, including any hand edits. Normally Ed & Sara edit
# people-roster.xlsx directly and it is the source of truth; re-seed only after a
# structural change to build_people.py's roster.
#
#   python3 scripts/seed-roster-xlsx.py
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

TSV = 'data/people.tsv'
OUT = 'data/people-roster.xlsx'
FOREST = "16342A"; YELLOW = "FFF3C4"; SAGE = "E7EADF"
arial = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="D3D9C8"); border = Border(left=thin, right=thin, top=thin, bottom=thin)

rows = list(csv.DictReader(open(TSV, encoding='utf-8'), delimiter='\t'))

wb = Workbook(); ws = wb.active; ws.title = "Roster"
labels = ["Name", "Role", "Start year", "End year", "Current position", "Website", "LinkedIn"]
ws.append(labels)
for c in range(1, 8):
    cell = ws.cell(1, c)
    cell.font = arial(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=FOREST)
    cell.alignment = Alignment(horizontal="left", vertical="center"); cell.border = border
ws.row_dimensions[1].height = 22
ws.cell(1, 2).comment = Comment("Choose from dropdown: pi, postdoc, grad, staff, undergrad, visiting", "guide")
ws.cell(1, 3).comment = Comment("4-digit year they joined (required to appear on the timeline)", "guide")
ws.cell(1, 4).comment = Comment("4-digit year they left, or 'present' (blank = present)", "guide")
ws.cell(1, 6).comment = Comment("Full URL, e.g. https://example.edu/person", "guide")
ws.cell(1, 7).comment = Comment("Full LinkedIn URL, e.g. https://www.linkedin.com/in/username", "guide")


def year(v):
    s = str(v or "").strip()
    if s.isdigit():
        return int(s)
    return s or None  # 'present' stays text; blank -> None


for r in rows:
    ws.append([r.get("name", ""), r.get("role", ""), year(r.get("start_year")),
               year(r.get("end_year")), r.get("current_position", "") or None, None, None])
nrows = ws.max_row
for row in range(2, nrows + 1):
    for col in range(1, 8):
        cell = ws.cell(row, col)
        cell.font = arial(size=10); cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=(col == 5))
        if col in (3, 4) and isinstance(cell.value, int):
            cell.number_format = '0'  # plain integer, no thousands comma
    sc = ws.cell(row, 3)
    if sc.value in (None, ""):
        sc.fill = PatternFill("solid", fgColor=YELLOW)

dv = DataValidation(type="list", formula1='"pi,postdoc,grad,staff,undergrad,visiting"', allow_blank=False)
dv.error = "Pick one of: pi, postdoc, grad, staff, undergrad, visiting"; dv.errorTitle = "Invalid role"
dv.prompt = "Choose the person's role"; dv.promptTitle = "Role"
ws.add_data_validation(dv); dv.add("B2:B500")
for col, w in {"A": 26, "B": 12, "C": 11, "D": 11, "E": 48, "F": 34, "G": 34}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:G{nrows}"

info = wb.create_sheet("Instructions"); info.column_dimensions["A"].width = 100
lines = [
    ("Lab roster — how to fill this out", True, FOREST, "FFFFFF", 13),
    ("", False, None, None, 10),
    ("This sheet feeds the People timeline on the lab website. Edit the Roster tab, then run scripts/generate-people.py (or hand it to the web team).", False, None, None, 10),
    ("", False, None, None, 10),
    ("Columns", True, SAGE, "16342A", 11),
    ("• Name — full name.", False, None, None, 10),
    ("• Role — pick from the dropdown: pi, postdoc, grad, staff, undergrad, visiting.", False, None, None, 10),
    ("• Start year — 4-digit year they joined. REQUIRED — people without a start year can't be placed on the timeline (yellow cells).", False, None, None, 10),
    ("• End year — 4-digit year they left, or 'present'. Blank counts as 'present'.", False, None, None, 10),
    ("• Current position — where they are now (commas/semicolons are fine).", False, None, None, 10),
    ("• Website — full URL (optional); shows as a globe icon by their name.", False, None, None, 10),
    ("• LinkedIn — full LinkedIn URL (optional); shows as a LinkedIn icon by their name.", False, None, None, 10),
    ("", False, None, None, 10),
    ("Notes", True, SAGE, "16342A", 11),
    ("• Many years were inferred from co-authorship, so they're approximate — especially common surnames (Miller, Li, Zhang, Wang, Chen). Verify and fix.", False, None, None, 10),
    ("• Yellow start-year cells are missing and need a year — mostly staff, undergrads, and short-term visitors.", False, None, None, 10),
    ("• Add new people as new rows at the bottom; order doesn't matter (the site sorts by start year).", False, None, None, 10),
]
for i, (txt, bold, fill, color, size) in enumerate(lines, 1):
    c = info.cell(i, 1, txt)
    c.font = arial(bold=bold, size=size, color=(color or "16201B"))
    c.alignment = Alignment(wrap_text=True, vertical="center")
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if bold:
        info.row_dimensions[i].height = 20

wb.save(OUT)
missing = sum(1 for r in rows if not str(r.get("start_year", "")).strip())
print(f"Wrote {OUT}: {len(rows)} people ({missing} missing start years highlighted)")
